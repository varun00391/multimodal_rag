from __future__ import annotations

import io
import json
from pathlib import Path

from PIL import Image

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".html", ".log", ".py", ".xml", ".yml", ".yaml"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff"}
VIDEO_EXT = {".mp4", ".mov", ".webm", ".mkv", ".avi"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac"}


def detect_modality(filename: str, mime: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf" or "pdf" in mime:
        return "pdf"
    if ext in IMAGE_EXT or mime.startswith("image/"):
        return "image"
    if ext in VIDEO_EXT or mime.startswith("video/"):
        return "video"
    if ext in AUDIO_EXT or mime.startswith("audio/"):
        return "audio"
    if ext == ".docx":
        return "docx"
    if ext == ".pptx":
        return "pptx"
    if ext in {".xlsx", ".xls"}:
        return "spreadsheet"
    return "text"


def chunk_text(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    cleaned = " ".join(text.split())
    if not cleaned:
        return []
    chunks: list[str] = []
    i = 0
    while i < len(cleaned):
        chunks.append(cleaned[i : i + size])
        i += max(size - overlap, 1)
    return chunks[:80]


def extract_text(filename: str, mime: str, data: bytes) -> tuple[str, dict]:
    ext = Path(filename).suffix.lower()
    meta: dict = {"bytes": len(data)}
    try:
        if ext == ".pdf" or "pdf" in mime:
            return _pdf(data), meta
        if ext == ".docx":
            return _docx(data), meta
        if ext == ".pptx":
            return _pptx(data), meta
        if ext in {".xlsx", ".xls"}:
            return _xlsx(data), meta
        if ext in IMAGE_EXT or mime.startswith("image/"):
            return _image(filename, data, meta)
        if ext in VIDEO_EXT or mime.startswith("video/"):
            return (
                f"Video file '{filename}' ingested for multimodal retrieval. "
                "No transcript was attached. Ask about this asset by filename, meeting, or product context.",
                {**meta, "kind": "video"},
            )
        if ext in AUDIO_EXT or mime.startswith("audio/"):
            return (
                f"Audio file '{filename}' ingested. No speech-to-text run in this prototype. "
                "The file is searchable by name and any surrounding notes.",
                {**meta, "kind": "audio"},
            )
        if ext in TEXT_EXT or mime.startswith("text/") or "json" in mime:
            return data.decode("utf-8", errors="ignore"), meta
    except Exception as exc:
        return f"Could not fully parse {filename}: {exc}", {"error": str(exc)}
    return data.decode("utf-8", errors="ignore") or f"Binary file {filename} stored for retrieval.", meta


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages[:40]):
        text = page.extract_text() or ""
        pages.append(f"[Page {i + 1}]\n{text}")
    return "\n\n".join(pages).strip() or "PDF contained no extractable text (may be scanned)."


def _docx(data: bytes) -> str:
    from docx import Document as Docx

    doc = Docx(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text).strip() or "Empty Word document."


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        bits = [shape.text for shape in slide.shapes if getattr(shape, "text", None)]
        slides.append(f"[Slide {i}] " + " ".join(bits))
    return "\n".join(slides).strip() or "Empty presentation."


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets[:8]:
        lines.append(f"[Sheet {sheet.title}]")
        for row in sheet.iter_rows(max_row=80, max_col=12, values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines).strip() or "Empty spreadsheet."


def _image(filename: str, data: bytes, meta: dict) -> tuple[str, dict]:
    img = Image.open(io.BytesIO(data))
    meta.update({"width": img.width, "height": img.height, "mode": img.mode})
    text = (
        f"Image '{filename}' ({img.width}x{img.height}, {img.mode}). "
        "Visual document stored for multimodal RAG. Describe what is in the photo when asking questions, "
        "or query by filename, location, SKU, or screenshot UI labels."
    )
    return text, meta


def extra_dump(meta: dict) -> str:
    return json.dumps(meta, default=str)
