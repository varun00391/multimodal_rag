from __future__ import annotations

from pathlib import Path

from extraction.elements import picture_element, table_element, text_element
from extraction.models.cdr import Diagnostic, Unit
from extraction.models.report import ExtractionReport, RouteRecord

VERSION = "1.0"


def extract_docx(
    path: Path,
    assets_dir: Path,
    report: ExtractionReport,
    diagnostics: list[Diagnostic],
) -> list[Unit]:
    del diagnostics
    from docx import Document

    document = Document(path)
    elements = []
    order = 1
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        element_type = "heading" if "heading" in style else "paragraph"
        elements.append(
            text_element(
                unit_type="file",
                unit_index=1,
                order=order,
                text=text,
                extractor="python-docx",
                version=VERSION,
                element_type=element_type,
            )
        )
        order += 1
    for table in document.tables:
        raw = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
        if not raw:
            continue
        elements.append(
            table_element(
                unit_type="file",
                unit_index=1,
                order=order,
                columns=raw[0],
                rows=raw[1:],
                extractor="python-docx",
                version=VERSION,
            )
        )
        order += 1
    image_index = 1
    for rel in document.part.rels.values():
        if "image" not in getattr(rel, "reltype", ""):
            continue
        blob = rel.target_part.blob
        suffix = Path(rel.target_ref).suffix or ".png"
        dest = assets_dir / f"docx-image-{image_index}{suffix}"
        dest.write_bytes(blob)
        elements.append(
            picture_element(
                unit_type="file",
                unit_index=1,
                order=order,
                asset_path=dest.name,
                extractor="python-docx",
                version=VERSION,
            )
        )
        image_index += 1
        order += 1
    report.routes.append(RouteRecord(unit_type="file", index=1, extractor="python-docx"))
    return [
        Unit(
            unit_type="file",
            index=1,
            primary_route="python-docx",
            elements=elements,
            error=None if elements else "DOCX_EMPTY",
        )
    ]


def extract_pptx(
    path: Path,
    assets_dir: Path,
    report: ExtractionReport,
    diagnostics: list[Diagnostic],
) -> list[Unit]:
    del diagnostics
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    units: list[Unit] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        elements = []
        order = 1
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text and paragraph.text.strip()
                ).strip()
                if text:
                    elements.append(
                        text_element(
                            unit_type="slide",
                            unit_index=slide_index,
                            order=order,
                            text=text,
                            extractor="python-pptx",
                            version=VERSION,
                        )
                    )
                    order += 1
            if shape.has_table:
                table = shape.table
                raw = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                if raw:
                    elements.append(
                        table_element(
                            unit_type="slide",
                            unit_index=slide_index,
                            order=order,
                            columns=raw[0],
                            rows=raw[1:],
                            extractor="python-pptx",
                            version=VERSION,
                        )
                    )
                    order += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                suffix = f".{image.ext}" if image.ext else ".png"
                dest = assets_dir / f"slide-{slide_index}-image-{order}{suffix}"
                dest.write_bytes(image.blob)
                elements.append(
                    picture_element(
                        unit_type="slide",
                        unit_index=slide_index,
                        order=order,
                        asset_path=dest.name,
                        extractor="python-pptx",
                        version=VERSION,
                    )
                )
                order += 1
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        if notes:
            elements.append(
                text_element(
                    unit_type="slide",
                    unit_index=slide_index,
                    order=order,
                    text=notes,
                    extractor="python-pptx",
                    version=VERSION,
                    element_type="notes",
                )
            )
        report.routes.append(RouteRecord(unit_type="slide", index=slide_index, extractor="python-pptx"))
        units.append(
            Unit(
                unit_type="slide",
                index=slide_index,
                primary_route="python-pptx",
                elements=elements,
                error=None if elements else "PPTX_SLIDE_EMPTY",
            )
        )
    return units
