from __future__ import annotations

from pathlib import Path

from extraction.detect import CSV, DOCX, PDF, detect_media_type, family_for
from extraction.errors import ExtractionError
from extraction.extractors.archive import unpack_zip
from extraction.models.cdr import CanonicalDocument
from extraction.pipeline import run_pipeline


def test_detect_csv(tmp_path: Path) -> None:
    path = tmp_path / "rows.csv"
    path.write_text("name,age\nAda,36\n", encoding="utf-8")
    assert detect_media_type(path) == CSV
    assert family_for(CSV) == "tabular"


def test_detect_pdf_header(tmp_path: Path) -> None:
    path = tmp_path / "doc.pdf"
    path.write_bytes(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    assert detect_media_type(path) == PDF
    assert family_for(PDF) == "pdf"


def test_detect_docx(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("Hello")
    document.save(path)
    assert detect_media_type(path, "note.docx") == DOCX
    assert family_for(DOCX) == "office"


def test_csv_pipeline(settings) -> None:
    source = settings.workspace / "sample.csv"
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_text("city,country\nParis,France\nTokyo,Japan\n", encoding="utf-8")
    document = run_pipeline(settings, source, "sample.csv", job_id="job-csv")
    assert document.status == "completed"
    assert document.unit_count == 1
    table = document.units[0].elements[0].table
    assert table is not None
    assert table.columns == ["city", "country"]
    assert table.rows[0] == ["Paris", "France"]
    CanonicalDocument.model_validate(document.model_dump())


def test_docx_pipeline(settings) -> None:
    from docx import Document

    source = settings.workspace / "note.docx"
    document = Document()
    document.add_heading("Policy", level=1)
    document.add_paragraph("Employees receive twenty days of leave.")
    document.save(source)
    result = run_pipeline(settings, source, "note.docx", job_id="job-docx")
    assert result.units[0].primary_route == "python-docx"
    texts = [element.text for element in result.units[0].elements if element.text]
    assert any("twenty days" in text for text in texts)


def test_dense_pdf_pipeline(settings) -> None:
    import pymupdf

    source = settings.workspace / "dense.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Leave policy covers vacation, sick leave, and holidays for all staff.")
    doc.save(source)
    doc.close()
    document = run_pipeline(settings, source, "dense.pdf", job_id="job-pdf")
    assert document.units[0].primary_route == "pymupdf"
    assert any(element.text for element in document.units[0].elements)


def test_pdf_table_bbox_accepts_tuple() -> None:
    from extraction.extractors.pdf import _bbox
    from extraction.models.cdr import BBox

    box = _bbox((10.0, 20.0, 110.0, 80.0))
    assert box == BBox(left=10.0, top=20.0, right=110.0, bottom=80.0)


def test_pdf_with_table_pipeline(settings) -> None:
    import pymupdf

    source = settings.workspace / "table.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 48), "Leave policy covers vacation, sick leave, and holidays for all staff.")
    x0, y0, x1, y1 = 72, 100, 400, 180
    page.draw_rect(pymupdf.Rect(x0, y0, x1, y1))
    page.draw_line(pymupdf.Point(236, y0), pymupdf.Point(236, y1))
    page.draw_line(pymupdf.Point(x0, 140), pymupdf.Point(x1, 140))
    page.insert_text((80, 122), "Name")
    page.insert_text((250, 122), "Role")
    page.insert_text((80, 162), "Ada")
    page.insert_text((250, 162), "Engineer")
    doc.save(source)
    doc.close()
    document = run_pipeline(settings, source, "table.pdf", job_id="job-pdf-table")
    assert document.status in {"completed", "completed_with_warnings"}
    tables = [element.table for element in document.units[0].elements if element.table]
    assert tables
    assert tables[0].columns == ["Name", "Role"]
    assert tables[0].rows[0] == ["Ada", "Engineer"]
    assert any(element.bbox is not None for element in document.units[0].elements if element.table)


def test_zip_rejects_unsafe_path(settings, tmp_path: Path) -> None:
    import zipfile

    archive = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("../escape.txt", "nope")
    try:
        unpack_zip(settings, archive, tmp_path / "out")
        raised = False
    except ExtractionError as exc:
        raised = exc.code == "ZIP_UNSAFE_PATH"
    assert raised
